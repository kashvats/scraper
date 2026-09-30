import csv, io, json, socket, time
from concurrent.futures import ThreadPoolExecutor
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import auth, settings, store
from app import app
from scraper import Config,crawl
from egress_proxy import resolve_public
PASSWORD='correct-horse-battery-testing'
HASH=auth.password_hash(PASSWORD)
@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'DATA',tmp_path);monkeypatch.setattr(settings,'DB',tmp_path/'test.sqlite3');monkeypatch.setattr(settings,'PUBLIC_ORIGIN','http://testserver')
    store.migrate()
    with store.connection(True) as db:db.executemany('INSERT INTO users VALUES (?,?,?,0)',[('alice','alice@example.com',HASH),('bob','bob@example.com',HASH)])
    with TestClient(app) as c:yield c

def login(c,email='alice@example.com'):
    r=c.post('/api/auth/login',json={'email':email,'password':PASSWORD});assert r.status_code==200,r.text
    c.headers['X-CSRF-Token']=r.json()['csrf'];return r

def config(**kw):return Config(url='https://shop.example/',fields=[{'name':'title','selector':'h1'}],delay_ms=0,retries=0,**kw)
class FakeBrowser:
    def __init__(self,c):self.url=''
    def __enter__(self):return self
    def __exit__(self,*a):pass
    def visit(self,url,leaf):self.url=url;return url
    def extract(self,c):
        if self.url.endswith('/'):return {'rows':[],'links':['/category','/category#duplicate'],'next':[]}
        if self.url.endswith('/category'):return {'rows':[],'links':['/product/1'],'next':['/category?page=2']}
        if 'page=2' in self.url:return {'rows':[],'links':['/product/1','/product/2'],'next':['/category']}
        return {'rows':[{'title':self.url.rsplit('/',1)[-1]}],'links':[],'next':[]}
def test_traversal_limits():
    rows,errors,pages,limited=crawl(config(link_levels=['a','a']),lambda *a:None,lambda:False,FakeBrowser)
    assert [r['title'] for r in rows]==['1','2'] and pages==5 and not errors and not limited
    rows,_,pages,limited=crawl(config(link_levels=['a','a'],max_pages=2),lambda *a:None,lambda:False,FakeBrowser)
    assert not rows and pages==2 and limited

def test_recovery_and_fencing(client):
    jid=store.create_job('alice',config(link_levels=['a','a']).model_dump());j=store.claim('w1');sent=0
    def persist(state,rows,errors,pages,url):
        nonlocal sent
        store.checkpoint(jid,j['lease'],state,rows[sent:],errors,pages,url);sent=len(rows)
        if pages==3:raise RuntimeError('crash after commit')
    with pytest.raises(RuntimeError):crawl(config(link_levels=['a','a']),lambda *a:None,lambda:False,FakeBrowser,on_checkpoint=persist)
    with store.connection(True) as db:db.execute('UPDATE jobs SET heartbeat=? WHERE id=?',(time.time()-120,jid))
    recovered=store.claim('w2');assert recovered['lease']!=j['lease'];sent=0
    def resumed(state,rows,errors,pages,url):
        nonlocal sent
        store.checkpoint(jid,recovered['lease'],state,rows[sent:],errors,pages,url);sent=len(rows)
    crawl(config(link_levels=['a','a']),lambda *a:None,lambda:False,FakeBrowser,checkpoint=recovered['checkpoint'],on_checkpoint=resumed)
    with store.connection() as db:rows=[json.loads(r[0]) for r in db.execute('SELECT payload FROM results WHERE job_id=? ORDER BY seq',(jid,))]
    assert [r['title'] for r in rows]==['1','2']
    with pytest.raises(RuntimeError):store.checkpoint(jid,j['lease'],{},[],[],0,'')
def test_atomic_claim(client):
    store.create_job('alice',config().model_dump())
    with ThreadPoolExecutor(4) as pool:claims=list(pool.map(store.claim,['w1','w2','w3','w4']))
    assert sum(c is not None for c in claims)==1

def test_auth_csrf_logout(client):
    assert client.get('/api/jobs').status_code==401
    r=login(client);assert 'HttpOnly' in r.headers['set-cookie'] and 'SameSite=strict' in r.headers['set-cookie']
    csrf=client.headers.pop('X-CSRF-Token');assert client.post('/api/jobs',json=config().model_dump()).status_code==403
    client.headers['X-CSRF-Token']=csrf
    assert client.post('/api/jobs',json=config().model_dump(),headers={'Origin':'https://evil.example'}).status_code==403
    assert client.post('/api/auth/logout').status_code==200
    assert client.get('/api/jobs').status_code==401

def test_ownership(client):
    login(client);jid=client.post('/api/jobs',json=config().model_dump()).json()['id']
    tid=client.post('/api/templates',json={'name':'Private','config':config().model_dump()}).json()['id']
    login(client,'bob@example.com')
    assert client.get('/api/jobs').json()==[] and client.get('/api/templates').json()==[]
    for path in [f'/api/jobs/{jid}',f'/api/jobs/{jid}/export/csv']:assert client.get(path).status_code==404
    for suffix in ('cancel','resume'):assert client.post(f'/api/jobs/{jid}/{suffix}').status_code==404
    client.delete('/api/templates/'+tid);login(client);assert len(client.get('/api/templates').json())==1

def test_exports(client):
    login(client);jid=store.create_job('alice',config().model_dump());j=store.claim('w')
    store.checkpoint(jid,j['lease'],{},[{'title':'=1+2'},{'title':'₹499, "quoted"\nline'},{'title':0}],[],1,'https://shop.example/')
    store.finish(jid,j['lease'],'completed')
    rows=list(csv.reader(io.StringIO(client.get(f'/api/jobs/{jid}/export/csv').content.decode('utf-8-sig'))))
    assert rows[1][0]=="'=1+2" and rows[2][0]=='₹499, "quoted"\nline' and rows[3][0]=='0'
    wb=load_workbook(io.BytesIO(client.get(f'/api/jobs/{jid}/export/xlsx').content));assert wb.active['A2'].data_type=='s' and wb.active.max_row==4

def test_cancel_resume_delete(client):
    login(client);jid=client.post('/api/jobs',json=config().model_dump()).json()['id']
    assert client.delete('/api/jobs/'+jid).status_code==409
    client.post(f'/api/jobs/{jid}/cancel');assert client.get('/api/jobs/'+jid).json()['status']=='cancelled'
    assert client.post(f'/api/jobs/{jid}/resume').status_code==202
    assert client.post(f'/api/jobs/{jid}/resume').status_code==409
    client.post(f'/api/jobs/{jid}/cancel');assert client.delete('/api/jobs/'+jid).status_code==200
    assert store.get_job(jid) is None

def test_validation_headers(client):
    login(client)
    assert client.post('/api/jobs',json={'url':'file:///etc/passwd','fields':[{'name':'x','selector':'h1'}]}).status_code==422
    assert client.post('/api/picker',json={'url':'file:///etc/passwd'}).status_code==422
    assert client.post('/api/picker/expired/action',json={'kind':'select','x':1001}).status_code==422
    r=client.get('/api/jobs');assert r.headers['cache-control']=='no-store' and r.headers['x-content-type-options']=='nosniff'
    assert client.post('/api/auth/login',content=b'x'*(1024*1024+1)).status_code==413

def test_picker_ownership(client,monkeypatch):
    import picker
    class Fake:owner='alice'
    monkeypatch.setitem(picker.sessions,'private',Fake());login(client,'bob@example.com')
    assert client.post('/api/picker/private/action',json={'kind':'refresh'}).status_code==404
    assert client.delete('/api/picker/private').status_code==404
    picker.sessions.pop('private')

@pytest.mark.parametrize('ip',['127.0.0.1','10.0.0.1','169.254.169.254','192.168.1.1','::1','fc00::1','0.0.0.0','224.0.0.1'])
def test_private_network_blocked(monkeypatch,ip):
    monkeypatch.setattr(socket,'getaddrinfo',lambda *a,**k:[(socket.AF_INET,socket.SOCK_STREAM,6,'',(ip,443))])
    with pytest.raises(ValueError):resolve_public('attacker.example',443)
def test_dns_mixed_records_ports(monkeypatch):
    monkeypatch.setattr(socket,'getaddrinfo',lambda *a,**k:[(socket.AF_INET,socket.SOCK_STREAM,6,'',('1.1.1.1',443)),(socket.AF_INET,socket.SOCK_STREAM,6,'',('127.0.0.1',443))])
    with pytest.raises(ValueError):resolve_public('mixed.example',443)
    with pytest.raises(ValueError):resolve_public('public.example',22)
def test_readiness_retention(client):
    assert client.get('/health/ready').status_code==503
    store.claim('healthy');assert client.get('/health/ready').status_code==200
    jid=store.create_job('alice',config().model_dump())
    with store.connection(True) as db:db.execute("UPDATE jobs SET status='completed',updated=0 WHERE id=?",(jid,))
    store.cleanup();assert store.get_job(jid) is None

def test_login_throttle(client):
    for _ in range(20):assert client.post('/api/auth/login',json={'email':'alice@example.com','password':'wrong'}).status_code==401
    assert client.post('/api/auth/login',json={'email':'alice@example.com','password':'wrong'}).status_code==429

def test_expired_session_and_disabled_user(client):
    login(client)
    with store.connection(True) as db:db.execute('UPDATE sessions SET expires=0')
    assert client.get('/api/jobs').status_code==401
    login(client)
    with store.connection(True) as db:db.execute("UPDATE users SET disabled=1 WHERE id='alice'")
    assert client.get('/api/jobs').status_code==401

def test_stale_picker_revision_rejected():
    from picker import PageSession,Action
    session=PageSession(None,'test');session.revision=2
    with pytest.raises(ValueError,match='view changed'):session.act(Action(kind='select',revision=1,x=10,y=10))
